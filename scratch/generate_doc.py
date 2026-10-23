import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_document():
    doc = Document()
    
    # Page setup - A4
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)
        
    # Styles
    COLOR_PRIMARY = RGBColor(30, 58, 138)     # Navy #1E3A8A
    COLOR_SECONDARY = RGBColor(37, 99, 235)  # Blue #2563EB
    COLOR_DARK = RGBColor(30, 41, 59)        # Slate #1E293B
    COLOR_MUTED = RGBColor(100, 116, 139)    # Muted #64748B
    
    # Header Banner
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    title_p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_cat = title_p.add_run("CUỘC THI TOÀN QUỐC: AI FOR LIFE 2026\n")
    run_cat.font.name = "Arial"
    run_cat.font.size = Pt(11)
    run_cat.font.bold = True
    run_cat.font.color.rgb = COLOR_SECONDARY
    
    run_title = title_p.add_run("HỒ SƠ Ý TƯỞNG & THUYẾT MINH DỰ ÁN")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY
    
    # Subtitle / App Name Box
    app_table = doc.add_table(rows=1, cols=1)
    app_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    app_cell = app_table.cell(0, 0)
    set_cell_background(app_cell, "F1F5F9")
    set_cell_margins(app_cell, top=140, bottom=140, left=200, right=200)
    app_p = app_cell.paragraphs[0]
    app_p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    app_p.paragraph_format.space_after = Pt(2)
    
    r_app = app_p.add_run("Tên Ứng Dụng: DocuGraph AI\n")
    r_app.font.name = "Arial"
    r_app.font.size = Pt(15)
    r_app.font.bold = True
    r_app.font.color.rgb = COLOR_PRIMARY
    
    r_tag = app_p.add_run("Tagline: \"Biến hàng ngàn trang tài liệu phức tạp thành Đồ thị Tri thức sống động trong tầm tay.\"")
    r_tag.font.name = "Arial"
    r_tag.font.size = Pt(11)
    r_tag.font.italic = True
    r_tag.font.color.rgb = COLOR_DARK
    
    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    
    def add_section_header(number_str, title_str):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        
        r_num = h.add_run(f"{number_str}. ")
        r_num.font.name = "Arial"
        r_num.font.size = Pt(13)
        r_num.font.bold = True
        r_num.font.color.rgb = COLOR_SECONDARY
        
        r_txt = h.add_run(title_str)
        r_txt.font.name = "Arial"
        r_txt.font.size = Pt(13)
        r_txt.font.bold = True
        r_txt.font.color.rgb = COLOR_PRIMARY
        return h

    def add_bullet(p, bold_prefix, text):
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.left_indent = Inches(0.2)
        r_bullet = p.add_run("•  ")
        r_bullet.font.name = "Arial"
        r_bullet.font.bold = True
        r_bullet.font.color.rgb = COLOR_SECONDARY
        
        if bold_prefix:
            r_bold = p.add_run(bold_prefix + " ")
            r_bold.font.name = "Arial"
            r_bold.font.size = Pt(10.5)
            r_bold.font.bold = True
            r_bold.font.color.rgb = COLOR_DARK
            
        r_main = p.add_run(text)
        r_main.font.name = "Arial"
        r_main.font.size = Pt(10.5)
        r_main.font.color.rgb = COLOR_DARK

    def add_paragraph(text, space_after=6):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(10.5)
        r.font.color.rgb = COLOR_DARK
        return p

    # 1. VẤN ĐỀ ĐỜI THƯỜNG
    add_section_header("1", "Vấn Đề Đời Thường Cần Giải Quyết")
    add_paragraph("Trong thời đại số hóa, con người đối mặt với nghịch lý \"ngộ độc dữ liệu nhưng lại thiếu hụt tri thức sâu sắc\" (Data-rich, Information-poor):")
    add_bullet(doc.add_paragraph(), "Dữ liệu phân mảnh & khó đọc:", "Hàng triệu trang sách giáo trình, nghiên cứu y khoa, văn bản quy phạm pháp luật, hợp đồng bảo hiểm tồn tại dưới dạng bản in scan mờ, giáo trình photo cũ, slide ảnh chụp hay chữ viết tay.")
    add_bullet(doc.add_paragraph(), "Đứt gãy liên kết ngữ cảnh:", "Người đọc chỉ có thể đọc tuần tự từng trang, rất khó xâu chuỗi và tìm ra mối tương quan giữa các khái niệm, điều khoản, hiện tượng nằm rải rác ở các chương hoặc các bộ tài liệu khác nhau.")
    add_bullet(doc.add_paragraph(), "Lãng phí thời gian trí tuệ:", "Mỗi sinh viên, nhà nghiên cứu, bác sĩ hay luật sư mất trung bình từ 1.5 – 3 giờ mỗi ngày chỉ để lướt tìm từ khóa, ghi chú thủ công và chắp vá thông tin.")

    # 2. ĐỐI TƯỢNG & HOÀN CẢNH
    add_section_header("2", "Đối Tượng Người Dùng & Hoàn Cảnh Xảy Ra Vấn Đề")
    add_bullet(doc.add_paragraph(), "Sinh viên, Giảng viên & Học viên Cao học:", "Đối mặt với hàng chục tài liệu giáo trình scan, tài liệu tham khảo bằng tiếng Việt và ngoại ngữ trong mùa thi, làm luận văn tốt nghiệp; gặp khó khăn khi liên kết kiến thức liên môn.")
    add_bullet(doc.add_paragraph(), "Bác sĩ, Dược sĩ & Chuyên gia Y tế:", "Cần tra cứu khẩn cấp phác đồ điều trị, cơ chế dược lý hoặc tương tác thuốc từ hàng ngàn trang dược thư để đưa ra quyết định lâm sàng chính xác cho bệnh nhân.")
    add_bullet(doc.add_paragraph(), "Luật sư, Chuyên viên Pháp lý & Doanh nghiệp:", "Cần đối chiếu hàng loạt văn bản quy phạm pháp luật, thông tư, nghị định sửa đổi bổ sung rải rác nhiều năm để thẩm định hợp đồng và tuân thủ pháp luật.")
    add_bullet(doc.add_paragraph(), "Nhân viên công sở & Khối Cơ quan:", "Cần tra cứu nhanh SOP (quy trình vận hành), biên bản cuộc họp, hợp đồng kinh tế đã được scan và lưu trữ trong kho tài liệu số.")

    # 3. VÌ SAO ĐÁNG ĐƯỢC GIẢI QUYẾT & BẤT TIỆN CỦA GIẢI PHÁP HIỆN TẠI
    add_section_header("3", "Tính Cấp Thiết & Hạn Chế Của Các Giải Pháp Hiện Tại")
    add_paragraph("Ý nghĩa thực tiễn: Giải quyết bài toán này giúp tiết kiệm hàng triệu giờ lao động trí tuệ, giảm thiểu sai sót nghiêm trọng do thiếu sót thông tin trong y tế, pháp lý, và bình đẳng hóa cơ hội tiếp cận tri thức chất lượng cao cho mọi người dân.")
    add_paragraph("Hạn chế của các giải pháp hiện nay:")
    add_bullet(doc.add_paragraph(), "Ctrl + F truyền thống:", "Chỉ tìm từ khóa chính xác từng ký tự; hoàn toàn bất lực trước file PDF scan ảnh và không thể hiểu ngữ nghĩa tương đồng.")
    add_bullet(doc.add_paragraph(), "ChatGPT / LLM phổ thông:", "Dễ bị ảo giác (hallucination) tạo thông tin sai sự thật; bị giới hạn cửa sổ ngữ cảnh (context window); không trích dẫn được trang tài liệu gốc.")
    add_bullet(doc.add_paragraph(), "RAG truyền thống (Vector Search thuần túy):", "Chỉ tìm kiếm từng đoạn văn bản rời rạc, hoàn toàn đứt gãy khả năng suy luận đa bước (multi-hop reasoning) khi câu trả lời đòi hỏi phải liên kết nhiều điều khoản ở các trang khác nhau.")

    # 4. GIẢI PHÁP CỦA ĐỘI THI
    add_section_header("4", "Ứng Dụng DocuGraph AI Làm Gì?")
    add_paragraph("DocuGraph AI là nền tảng Trợ lý Tri thức Đa tầng toàn diện, thực hiện 4 nhiệm vụ cốt lõi:")
    add_bullet(doc.add_paragraph(), "Tiếp nhận đa định dạng:", "Tự động đọc hiểu và bóc tách chữ từ tài liệu PDF số, PDF scan mờ, ảnh chụp tài liệu, slide bài giảng tiếng Việt và tiếng Anh.")
    add_bullet(doc.add_paragraph(), "Kiến tạo Đồ thị Tri thức (Knowledge Graph):", "Tự động nhận diện các thực thể (Khái niệm, Nhân vật, Điều luật, Triệu chứng,...) và mối liên hệ nhân - quả giữa chúng thành một mạng lưới mạng nhện sống động.")
    add_bullet(doc.add_paragraph(), "Hỏi đáp GraphRAG thông minh:", "Trả lời các câu hỏi tự nhiên bằng tiếng Việt, tổng hợp tri thức đa tầng, loại bỏ ảo giác và trích dẫn chuẩn xác số trang nguồn.")
    add_bullet(doc.add_paragraph(), "Tương tác trực quan 2D/3D:", "Cung cấp giao diện trực quan cho phép người dùng lướt chạm, mở rộng các nút thực thể để khám phá các mắt xích kiến thức chưa từng nhận ra.")

    # 5. CÔNG NGHỆ AI ĐƯỢC ỨNG DỤNG
    add_section_header("5", "Trí Tuệ Nhân Tạo (AI) Được Sử Dụng Như Thế Nào?")
    add_bullet(doc.add_paragraph(), "Computer Vision & OCR Đa luồng (PaddleOCR v4/v5):", "Xử lý nhận diện ký tự quang học chuyên sâu cho tiếng Việt có dấu, chạy song song đa tiến trình tối ưu hiệu năng ngay cả với tài liệu scan mờ hay nghiêng góc.")
    add_bullet(doc.add_paragraph(), "Trích xuất Thực thể & Quan hệ bằng Generative AI:", "Tận dụng Google Gemini 2.5 Flash với Structured Output (JSON Schema) để phân tích văn bản thành các Node và Edge có cấu trúc.")
    add_bullet(doc.add_paragraph(), "Semantic Embedding (Google Gemini Embedding):", "Chuyển hóa văn bản thành không gian vector 768 chiều để tìm kiếm tương đồng ngữ nghĩa chính xác tuyệt đối.")
    add_bullet(doc.add_paragraph(), "GraphRAG Hybrid Retrieval (Neo4j Vector & Cypher):", "Kết hợp tìm kiếm Vector (Vector Index) và duyệt đồ thị tri thức (Graph Traversal), giúp thực hiện suy luận đa bước (multi-hop reasoning) vượt trội hơn hẳn RAG truyền thống.")
    add_bullet(doc.add_paragraph(), "Fact-Grounded Generation & Streaming:", "Mô hình ngôn ngữ lớn tổng hợp câu trả lời dựa trên sự thật (fact-grounded), truyền phát câu trả lời thời gian thực (SSE) kèm trích dẫn số trang tài liệu.")

    # 6. CÁC BƯỚC SỬ DỤNG
    add_section_header("6", "Quy Trình Trải Nghiệm Của Người Dùng (User Journey)")
    add_paragraph("Quy trình 4 bước tinh gọn, thân thiện, không đòi hỏi kỹ năng công nghệ phức tạp:")
    add_bullet(doc.add_paragraph(), "Bước 1 - Tải tài liệu:", "Kéo thả tài liệu (PDF, scan ảnh, giáo trình) vào hệ thống.")
    add_bullet(doc.add_paragraph(), "Bước 2 - AI học tự động:", "Hệ thống tự động thực hiện OCR, tạo vector nhúng và dệt nên đồ thị tri thức trong tích tắc.")
    add_bullet(doc.add_paragraph(), "Bước 3 - Trò chuyện tự nhiên:", "Đặt câu hỏi bằng tiếng Việt tự nhiên (ví dụ: \"Tội danh này bị xử phạt ra sao theo quy định mới nhất?\").")
    add_bullet(doc.add_paragraph(), "Bước 4 - Xác thực & Khám phá:", "Xem câu trả lời tóm tắt, nhấp vào trích dẫn để nhảy tới trang gốc hoặc mở Đồ thị để khám phá mạng lưới khái niệm liên quan.")

    # 7. SƠ ĐỒ HOẠT ĐỘNG VÀ GIAO DIỆN
    add_section_header("7", "Sơ Đồ Kiến Trúc Hoạt Động & Giao Diện")
    add_paragraph("Kiến trúc phân tầng chuyên nghiệp, đảm bảo độ mở rộng và tính an toàn dữ liệu:")
    
    # Table for Architecture
    arch_table = doc.add_table(rows=5, cols=2)
    arch_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    arch_headers = ["Tầng Hệ Thống", "Công Nghệ & Thành Phần Chi Tiết"]
    for c_idx, h_text in enumerate(arch_headers):
        c = arch_table.cell(0, c_idx)
        set_cell_background(c, "1E3A8A")
        set_cell_margins(c, top=80, bottom=80, left=100, right=100)
        p = c.paragraphs[0]
        r = p.add_run(h_text)
        r.font.name = "Arial"
        r.font.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    arch_rows = [
        ("Giao Diện (Frontend)", "Next.js 15, React 19, TailwindCSS, Streaming Chat UI, Interactive Knowledge Graph Visualizer, Source Document Viewer."),
        ("Xử Lý Trung Tâm (Backend)", "FastAPI (Python 3.12+), REST API, Server-Sent Events (SSE), Hệ thống bảo mật JWT Authentication & RBAC."),
        ("Bộ Động Cơ AI", "PaddleOCR tiếng Việt, Google Gemini 2.5 Flash (LLM), Google Gemini Embedding 768-dim, LangChain Framework."),
        ("Lưu Trữ Dữ Liệu", "Neo4j 5.12 (Graph Database & Vector Store cosine similarity) song hành cùng PostgreSQL 16 (Dữ liệu quan hệ, User, Metadata).")
    ]
    for r_idx, (col1, col2) in enumerate(arch_rows, start=1):
        c1 = arch_table.cell(r_idx, 0)
        c2 = arch_table.cell(r_idx, 1)
        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        set_cell_background(c1, bg)
        set_cell_background(c2, bg)
        set_cell_margins(c1, top=60, bottom=60, left=100, right=100)
        set_cell_margins(c2, top=60, bottom=60, left=100, right=100)
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(col1)
        r1.font.name = "Arial"
        r1.font.bold = True
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = COLOR_DARK
        
        p2 = c2.paragraphs[0]
        r2 = p2.add_run(col2)
        r2.font.name = "Arial"
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = COLOR_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 8. KHẢ NĂNG PHÁT TRIỂN THỰC TẾ
    add_section_header("8", "Kế Hoạch Phát Triển Thực Tế & Mô Hình Bền Vững")
    add_bullet(doc.add_paragraph(), "Giai đoạn 1 - MVP Sẵn sàng (Hiện tại):", "Đã hoàn thiện hệ thống full-stack hoạt động trơn tru với OCR tiếng Việt, Neo4j GraphRAG và giao diện trực quan.")
    add_bullet(doc.add_paragraph(), "Giai đoạn 2 - Chuyên môn hóa ngành (3 - 6 tháng):", "Tích hợp cơ sở tri thức chuyên sâu ngành Y tế (Dược thư quốc gia) và Pháp luật Việt Nam; phát triển ứng dụng di động cho phép quét chụp sách bằng camera; tích hợp giọng nói (Voice AI).")
    add_bullet(doc.add_paragraph(), "Giai đoạn 3 - Doanh nghiệp & On-Premises (6 - 12 tháng):", "Đóng gói giải pháp triển khai cục bộ (Local LLM & Offline OCR) bảo mật 100% dữ liệu cho ngân hàng, bệnh viện, cơ quan tư pháp.")
    add_bullet(doc.add_paragraph(), "Mô hình kinh doanh:", "Mô hình Freemium phục vụ cộng đồng học sinh/sinh viên miễn phí; gói B2B SaaS trả phí linh hoạt cho trường đại học, bệnh viện và doanh nghiệp.")

    # 9. ĐIỂM KHÁC BIỆT & GIÁ TRỊ LỚN NHẤT
    add_section_header("9", "Điểm Khác Biệt & Giá Trị Lớn Nhất (UVP)")
    
    diff_table = doc.add_table(rows=5, cols=4)
    diff_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    diff_headers = ["Tiêu Chí So Sánh", "Chatbot Phổ Thông", "RAG Truyền Thống", "DocuGraph AI (Đội Thi)"]
    for c_idx, h_text in enumerate(diff_headers):
        c = diff_table.cell(0, c_idx)
        set_cell_background(c, "1E3A8A")
        set_cell_margins(c, top=80, bottom=80, left=80, right=80)
        p = c.paragraphs[0]
        r = p.add_run(h_text)
        r.font.name = "Arial"
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    diff_rows = [
        ("Nguồn dữ liệu", "Chỉ nhận text thuần", "Chỉ nhận text số hóa", "Đa định dạng (Scan mờ, ảnh chụp, sách cũ)"),
        ("Tư duy xâu chuỗi", "Kém, dễ lạc đề", "Tách đoạn rời rạc", "Suy luận đa tầng qua Đồ thị Tri thức"),
        ("Độ tin cậy", "Hay ảo giác (Hallucination)", "Khá, dẫn chứng chung", "100% minh bạch, dẫn chứng số trang"),
        ("Trực quan hóa", "Chỉ có khung chat chữ", "Chỉ có khung chat chữ", "Đồ thị tương tác 2D/3D sống động")
    ]
    for r_idx, (c1_txt, c2_txt, c3_txt, c4_txt) in enumerate(diff_rows, start=1):
        bg = "F1F5F9" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, txt in enumerate([c1_txt, c2_txt, c3_txt, c4_txt]):
            c = diff_table.cell(r_idx, c_idx)
            set_cell_background(c, bg)
            set_cell_margins(c, top=60, bottom=60, left=80, right=80)
            p = c.paragraphs[0]
            r = p.add_run(txt)
            r.font.name = "Arial"
            r.font.size = Pt(9)
            if c_idx == 3:
                r.font.bold = True
                r.font.color.rgb = COLOR_SECONDARY
            elif c_idx == 0:
                r.font.bold = True
                r.font.color.rgb = COLOR_DARK
            else:
                r.font.color.rgb = COLOR_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # CONCLUSION BOX
    concl_table = doc.add_table(rows=1, cols=1)
    concl_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    concl_cell = concl_table.cell(0, 0)
    set_cell_background(concl_cell, "EFF6FF")
    set_cell_margins(concl_cell, top=140, bottom=140, left=180, right=180)
    concl_p = concl_cell.paragraphs[0]
    concl_p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    r_c_bold = concl_p.add_run("THÔNG ĐIỆP GỬI ĐẾN BAN GIÁM KHẢO AI FOR LIFE:\n")
    r_c_bold.font.name = "Arial"
    r_c_bold.font.size = Pt(11)
    r_c_bold.font.bold = True
    r_c_bold.font.color.rgb = COLOR_PRIMARY
    
    r_c_text = concl_p.add_run("\"AI for Life không phải là thay thế con người, mà là trang bị cho mỗi cá nhân một bộ não thứ hai — có khả năng ghi nhớ trọn vẹn và kết nối hàng vạn trang tri thức của nhân loại, giúp chúng ta đưa ra quyết định nhanh hơn, chuẩn xác hơn và sống một cuộc đời ý nghĩa hơn.\"")
    r_c_text.font.name = "Arial"
    r_c_text.font.size = Pt(10.5)
    r_c_text.font.italic = True
    r_c_text.font.color.rgb = COLOR_DARK

    # Save document
    output_path = r"c:\Users\nguye\PycharmProjects\THUCTAP\Ho_So_Du_Thi_AI_for_Life_DocuGraph_AI.docx"
    doc.save(output_path)
    print(f"Document successfully created at: {output_path}")

if __name__ == "__main__":
    create_document()
