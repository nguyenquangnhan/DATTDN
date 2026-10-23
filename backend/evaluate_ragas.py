"""
Đánh giá hệ thống DocuGraph RAG bằng thư viện RAGAS
=====================================================

Pipeline:
1. Trích xuất text từ tài liệu PDF (dùng PyMuPDF)
2. Dùng Gemini để sinh bộ câu hỏi + câu trả lời ground-truth
3. Chạy từng câu hỏi qua RAG pipeline (vector search + graph context + LLM)
4. Đánh giá bằng RAGAS metrics
5. Xuất kết quả ra file Excel

Chạy:
    cd backend
    uv run python evaluate_ragas.py
"""

import asyncio
import json
import logging
import os
import sys
from datetime import datetime
from pathlib import Path

# ── Fix Unicode encoding on Windows console ──
if sys.platform == "win32":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

# ── Ensure backend package is importable ──
sys.path.insert(0, str(Path(__file__).resolve().parent))
os.chdir(Path(__file__).resolve().parent)

# Load .env BEFORE any app imports
from dotenv import load_dotenv
load_dotenv(".env")

import fitz  # PyMuPDF
from google import genai
from google.genai import types as genai_types
from pydantic import BaseModel

from app.config import get_settings

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s  %(levelname)-8s  %(message)s",
)
logger = logging.getLogger("ragas_eval")
settings = get_settings()


# =====================================================================
# PHẦN 1: TRÍCH XUẤT TEXT TỪ TÀI LIỆU
# =====================================================================

def extract_pdf_text(pdf_path: str) -> str:
    """Trích xuất toàn bộ text từ PDF bằng PyMuPDF."""
    logger.info(f"📄 Đang đọc PDF: {pdf_path}")
    doc = fitz.open(pdf_path)
    pages_text = []
    for page_num in range(len(doc)):
        page = doc[page_num]
        text = page.get_text("text")
        if text.strip():
            pages_text.append(text.strip())
    doc.close()
    full_text = "\n\n".join(pages_text)
    logger.info(f"✅ Đã trích xuất {len(full_text):,} ký tự từ {len(pages_text)} trang")
    return full_text


# =====================================================================
# PHẦN 2: SINH BỘ CÂU HỎI & CÂU TRẢ LỜI GROUND-TRUTH
# =====================================================================

class QAPair(BaseModel):
    question: str
    ground_truth: str


class QADataset(BaseModel):
    pairs: list[QAPair]


def generate_qa_pairs(document_text: str, num_questions: int = 15) -> list[dict]:
    """Dùng Gemini để sinh bộ câu hỏi + câu trả lời ground-truth từ nội dung tài liệu."""
    logger.info(f"🧠 Đang sinh {num_questions} cặp câu hỏi-trả lời bằng Gemini...")

    client = genai.Client(api_key=settings.GEMINI_API_KEY)

    # Cắt text nếu quá dài (giữ ~30k ký tự)
    truncated_text = document_text[:30000]

    prompt = f"""Dựa vào nội dung tài liệu dưới đây, hãy sinh ra chính xác {num_questions} cặp câu hỏi và câu trả lời.

YÊU CẦU:
- Câu hỏi phải đa dạng: bao gồm câu hỏi chi tiết, câu hỏi tổng hợp, câu hỏi so sánh, câu hỏi giải thích
- Câu trả lời (ground_truth) phải chính xác dựa trên nội dung tài liệu
- Câu trả lời nên đủ chi tiết (2-5 câu)
- Viết bằng tiếng Việt
- Không đặt câu hỏi quá đơn giản hoặc yes/no

NỘI DUNG TÀI LIỆU:
{truncated_text}
"""

    response = client.models.generate_content(
        model=settings.CHAT_MODEL,
        contents=prompt,
        config=genai_types.GenerateContentConfig(
            response_mime_type="application/json",
            response_schema=QADataset,
            temperature=0.5,
        ),
    )

    raw = json.loads(response.text)
    pairs = raw.get("pairs", [])
    logger.info(f"✅ Đã sinh {len(pairs)} cặp Q&A")

    return [{"question": p["question"], "ground_truth": p["ground_truth"]} for p in pairs]


# =====================================================================
# PHẦN 3: CHẠY RAG PIPELINE CHO TỪNG CÂU HỎI
# =====================================================================

async def run_rag_for_questions(
    qa_pairs: list[dict],
    notebook_id: str,
) -> list[dict]:
    """Chạy RAG pipeline cho từng câu hỏi, thu thập answer + retrieved contexts."""

    from app.services.embedding_service import embed_query
    from app.services.neo4j_service import get_neo4j_driver
    from app.services.rag_service import get_graph_context, _get_chat_model, SYSTEM_PROMPT

    logger.info(f"🔍 Chạy RAG cho {len(qa_pairs)} câu hỏi trên notebook: {notebook_id}")

    results = []
    driver = get_neo4j_driver()
    chat_model = _get_chat_model()
    top_k = settings.RAG_TOP_K

    for idx, qa in enumerate(qa_pairs):
        question = qa["question"]
        logger.info(f"  [{idx+1}/{len(qa_pairs)}] {question[:60]}...")

        try:
            # 1. Embed query
            query_embedding = await embed_query(question)

            # 2. Vector search in Neo4j
            with driver.session() as session:
                result = session.run(
                    """
                    CALL db.index.vector.queryNodes('chunk_vector_index', $top_k, $query_vector)
                    YIELD node, score
                    MATCH (node)-[:PART_OF]->(s:Source)-[:BELONGS_TO]->(n:Notebook {id: $notebook_id})
                    RETURN 
                        node.id AS id,
                        node.content AS content,
                        s.name AS source_name,
                        score AS similarity
                    """,
                    {
                        "notebook_id": notebook_id,
                        "query_vector": query_embedding,
                        "top_k": top_k,
                    },
                )
                chunks = []
                for row in result:
                    chunks.append({
                        "id": row["id"],
                        "content": row["content"],
                        "source_name": row["source_name"],
                        "similarity": float(row["similarity"]),
                    })

            # 3. Build contexts
            contexts = [c["content"] for c in chunks]
            chunk_ids = [c["id"] for c in chunks]

            # 4. Get graph context
            graph_ctx = get_graph_context(chunk_ids)

            # 5. Build prompt and call LLM
            context_str = "\n\n---\n\n".join(
                [f"[{c['source_name']}]:\n{c['content']}" for c in chunks]
            )
            system_content = SYSTEM_PROMPT.format(
                context=context_str, graph_context=graph_ctx
            )
            messages = [("system", system_content), ("human", question)]
            response = await chat_model.ainvoke(messages)

            results.append({
                "question": question,
                "ground_truth": qa["ground_truth"],
                "answer": response.content,
                "contexts": contexts,
                "num_chunks_retrieved": len(chunks),
                "avg_similarity": (
                    sum(c["similarity"] for c in chunks) / len(chunks)
                    if chunks
                    else 0
                ),
                "has_graph_context": bool(graph_ctx),
            })

        except Exception as e:
            logger.error(f"  ❌ Lỗi câu {idx+1}: {e}")
            results.append({
                "question": question,
                "ground_truth": qa["ground_truth"],
                "answer": f"[ERROR] {e}",
                "contexts": [],
                "num_chunks_retrieved": 0,
                "avg_similarity": 0,
                "has_graph_context": False,
            })

    logger.info(f"✅ Hoàn thành RAG cho {len(results)} câu hỏi")
    return results


# =====================================================================
# PHẦN 4: ĐÁNH GIÁ BẰNG RAGAS METRICS
# =====================================================================

def evaluate_with_ragas(rag_results: list[dict]) -> dict:
    """Đánh giá kết quả RAG bằng thư viện RAGAS."""
    from ragas import evaluate
    from ragas.metrics import (
        faithfulness,
        answer_relevancy,
        context_precision,
        context_recall,
    )
    from datasets import Dataset

    logger.info("📊 Đang chạy RAGAS evaluation...")

    # Chuẩn bị dataset theo format RAGAS yêu cầu
    eval_data = {
        "question": [],
        "answer": [],
        "contexts": [],
        "ground_truth": [],
    }

    for r in rag_results:
        if r["answer"].startswith("[ERROR]"):
            continue
        eval_data["question"].append(r["question"])
        eval_data["answer"].append(r["answer"])
        eval_data["contexts"].append(r["contexts"] if r["contexts"] else [""])
        eval_data["ground_truth"].append(r["ground_truth"])

    if not eval_data["question"]:
        logger.error("❌ Không có kết quả hợp lệ để đánh giá!")
        return {}

    dataset = Dataset.from_dict(eval_data)
    logger.info(f"  Dataset: {len(dataset)} samples")

    # Chạy đánh giá RAGAS
    result = evaluate(
        dataset,
        metrics=[faithfulness, answer_relevancy, context_precision, context_recall],
    )

    logger.info(f"✅ RAGAS evaluation hoàn thành!")
    logger.info(f"  Kết quả tổng hợp: {result}")

    return result


# =====================================================================
# PHẦN 5: XUẤT KẾT QUẢ RA FILE EXCEL
# =====================================================================

def export_to_excel(
    rag_results: list[dict],
    ragas_scores: dict,
    output_path: str,
):
    """Xuất toàn bộ kết quả đánh giá ra file Excel với 3 sheet."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    logger.info(f"📁 Đang xuất kết quả ra Excel: {output_path}")

    wb = openpyxl.Workbook()

    # ── Styles ──
    header_font = Font(name="Arial", bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    good_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    warn_fill = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")
    bad_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    title_font = Font(name="Arial", bold=True, size=14, color="1E40AF")
    thin_border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )

    def apply_header(ws, row, num_cols):
        for col in range(1, num_cols + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

    def apply_border_and_wrap(ws, row, num_cols):
        for col in range(1, num_cols + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    def score_fill(score):
        if score >= 0.8:
            return good_fill
        elif score >= 0.5:
            return warn_fill
        else:
            return bad_fill

    # ================================================================
    # SHEET 1: Chi tiết từng câu hỏi
    # ================================================================
    ws1 = wb.active
    ws1.title = "Chi tiết Q&A"

    # Title
    ws1.merge_cells("A1:H1")
    title_cell = ws1["A1"]
    title_cell.value = "Kết quả đánh giá RAG - Chi tiết từng câu hỏi"
    title_cell.font = title_font
    title_cell.alignment = Alignment(horizontal="center")

    ws1.merge_cells("A2:H2")
    ws1["A2"].value = f"Thời gian đánh giá: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}"
    ws1["A2"].font = Font(italic=True, color="6B7280")
    ws1["A2"].alignment = Alignment(horizontal="center")

    # Headers
    headers = ["STT", "Câu hỏi", "Câu trả lời Ground-truth", "Câu trả lời RAG",
               "Contexts Retrieved", "Số chunks", "Avg Similarity", "Có Graph Context"]
    for col, h in enumerate(headers, 1):
        ws1.cell(row=4, column=col, value=h)
    apply_header(ws1, 4, len(headers))

    # Data
    for i, r in enumerate(rag_results):
        row = 5 + i
        ws1.cell(row=row, column=1, value=i + 1)
        ws1.cell(row=row, column=2, value=r["question"])
        ws1.cell(row=row, column=3, value=r["ground_truth"])
        ws1.cell(row=row, column=4, value=r["answer"])
        ws1.cell(row=row, column=5, value="\n---\n".join(r["contexts"][:3]) if r["contexts"] else "N/A")
        ws1.cell(row=row, column=6, value=r["num_chunks_retrieved"])
        ws1.cell(row=row, column=7, value=round(r["avg_similarity"], 4))
        ws1.cell(row=row, column=8, value="✅ Có" if r["has_graph_context"] else "❌ Không")
        apply_border_and_wrap(ws1, row, len(headers))

    # Column widths
    ws1.column_dimensions["A"].width = 6
    ws1.column_dimensions["B"].width = 40
    ws1.column_dimensions["C"].width = 45
    ws1.column_dimensions["D"].width = 50
    ws1.column_dimensions["E"].width = 50
    ws1.column_dimensions["F"].width = 12
    ws1.column_dimensions["G"].width = 15
    ws1.column_dimensions["H"].width = 16

    # ================================================================
    # SHEET 2: Điểm RAGAS
    # ================================================================
    ws2 = wb.create_sheet("Điểm RAGAS")

    ws2.merge_cells("A1:D1")
    title2 = ws2["A1"]
    title2.value = "Điểm đánh giá RAGAS Metrics"
    title2.font = title_font
    title2.alignment = Alignment(horizontal="center")

    # Aggregate scores table
    ragas_headers = ["Metric", "Điểm số (0-1)", "Đánh giá", "Mô tả"]
    for col, h in enumerate(ragas_headers, 1):
        ws2.cell(row=3, column=col, value=h)
    apply_header(ws2, 3, len(ragas_headers))

    # Metric descriptions
    metric_info = {
        "faithfulness": (
            "Faithfulness (Tính trung thực)",
            "Đo lường câu trả lời có bám sát ngữ cảnh truy xuất được không, không bịa thông tin (hallucination)",
        ),
        "answer_relevancy": (
            "Answer Relevancy (Độ liên quan câu trả lời)",
            "Đo lường câu trả lời có đúng trọng tâm câu hỏi không",
        ),
        "context_precision": (
            "Context Precision (Độ chính xác ngữ cảnh)",
            "Tỷ lệ chunks truy xuất được thực sự liên quan đến câu hỏi",
        ),
        "context_recall": (
            "Context Recall (Độ phủ ngữ cảnh)",
            "Khả năng truy xuất đủ thông tin cần thiết để trả lời câu hỏi",
        ),
    }

    row_idx = 4
    for metric_key, (display_name, desc) in metric_info.items():
        score = ragas_scores.get(metric_key, 0) if ragas_scores else 0
        if score >= 0.8:
            rating = "✅ Tốt"
        elif score >= 0.6:
            rating = "⚠️ Khá"
        elif score >= 0.4:
            rating = "⚠️ Trung bình"
        else:
            rating = "❌ Cần cải thiện"

        ws2.cell(row=row_idx, column=1, value=display_name)
        score_cell = ws2.cell(row=row_idx, column=2, value=round(score, 4) if score else "N/A")
        if isinstance(score, (int, float)):
            score_cell.fill = score_fill(score)
            score_cell.alignment = Alignment(horizontal="center")
        ws2.cell(row=row_idx, column=3, value=rating)
        ws2.cell(row=row_idx, column=4, value=desc)
        apply_border_and_wrap(ws2, row_idx, len(ragas_headers))
        row_idx += 1

    # Average score
    row_idx += 1
    ws2.cell(row=row_idx, column=1, value="ĐIỂM TRUNG BÌNH")
    ws2.cell(row=row_idx, column=1).font = Font(bold=True, size=12)
    valid_scores = [
        v for k, v in (ragas_scores or {}).items()
        if isinstance(v, (int, float)) and k in metric_info
    ]
    avg = sum(valid_scores) / len(valid_scores) if valid_scores else 0
    avg_cell = ws2.cell(row=row_idx, column=2, value=round(avg, 4))
    avg_cell.font = Font(bold=True, size=12)
    avg_cell.fill = score_fill(avg)
    avg_cell.alignment = Alignment(horizontal="center")

    ws2.column_dimensions["A"].width = 45
    ws2.column_dimensions["B"].width = 18
    ws2.column_dimensions["C"].width = 20
    ws2.column_dimensions["D"].width = 70

    # ================================================================
    # SHEET 3: Thông tin cấu hình
    # ================================================================
    ws3 = wb.create_sheet("Cấu hình hệ thống")

    ws3.merge_cells("A1:B1")
    title3 = ws3["A1"]
    title3.value = "Cấu hình hệ thống RAG"
    title3.font = title_font
    title3.alignment = Alignment(horizontal="center")

    config_headers = ["Tham số", "Giá trị"]
    for col, h in enumerate(config_headers, 1):
        ws3.cell(row=3, column=col, value=h)
    apply_header(ws3, 3, 2)

    config_data = [
        ("LLM Model", settings.CHAT_MODEL),
        ("Embedding Model", settings.EMBEDDING_MODEL),
        ("Embedding Dimensions", "768"),
        ("Chunk Size", str(settings.CHUNK_SIZE)),
        ("Chunk Overlap", str(settings.CHUNK_OVERLAP)),
        ("RAG Top-K", str(settings.RAG_TOP_K)),
        ("Temperature (LLM)", "0.3"),
        ("Max Output Tokens", "4096"),
        ("Vector Similarity", "Cosine"),
        ("Graph Database", "Neo4j 5.12"),
        ("Relational Database", "PostgreSQL 16"),
        ("OCR Engine", f"PaddleOCR (lang={settings.PADDLE_OCR_LANG})"),
        ("Số câu hỏi đánh giá", str(len(rag_results))),
        ("Thời gian đánh giá", datetime.now().strftime("%d/%m/%Y %H:%M:%S")),
    ]

    for i, (param, val) in enumerate(config_data):
        row = 4 + i
        ws3.cell(row=row, column=1, value=param)
        ws3.cell(row=row, column=2, value=val)
        apply_border_and_wrap(ws3, row, 2)

    ws3.column_dimensions["A"].width = 30
    ws3.column_dimensions["B"].width = 50

    # Save
    wb.save(output_path)
    logger.info(f"✅ Đã xuất kết quả ra: {output_path}")


# =====================================================================
# PHẦN 5b: XUẤT KẾT QUẢ KHÔNG CẦN RAGAS (fallback khi chưa cài ragas)
# =====================================================================

def evaluate_without_ragas(rag_results: list[dict]) -> dict:
    """Đánh giá thủ công khi chưa cài đặt thư viện RAGAS.

    Sử dụng Gemini để đánh giá từng câu dựa trên 4 tiêu chí RAGAS.
    Có retry logic để xử lý rate limit (429).
    """
    import time
    import re

    logger.info("📊 Đánh giá bằng Gemini (thay thế RAGAS library)...")

    client = genai.Client(api_key=settings.GEMINI_API_KEY)

    class EvalScores(BaseModel):
        faithfulness: float
        answer_relevancy: float
        context_precision: float
        context_recall: float

    all_scores = []
    valid_items = [(idx, r) for idx, r in enumerate(rag_results) if not r["answer"].startswith("[ERROR]")]

    for pos, (idx, r) in enumerate(valid_items):
        contexts_text = "\n---\n".join(r["contexts"][:3]) if r["contexts"] else "N/A"

        eval_prompt = f"""Bạn là chuyên gia đánh giá hệ thống RAG. Hãy chấm điểm từ 0.0 đến 1.0 cho các tiêu chí sau:

1. **faithfulness**: Câu trả lời có bám sát ngữ cảnh (contexts) không? Có bịa thông tin không?
2. **answer_relevancy**: Câu trả lời có đúng trọng tâm câu hỏi không?
3. **context_precision**: Các chunks truy xuất được có liên quan đến câu hỏi không?
4. **context_recall**: Các chunks truy xuất được có đủ thông tin để trả lời không? (so sánh với ground_truth)

Câu hỏi: {r['question']}
Ground-truth: {r['ground_truth']}
Câu trả lời RAG: {r['answer']}
Contexts retrieved:
{contexts_text}

Hãy trả về điểm số cho 4 tiêu chí."""

        # Retry loop for rate limiting
        max_retries = 5
        for attempt in range(max_retries):
            try:
                response = client.models.generate_content(
                    model=settings.CHAT_MODEL,
                    contents=eval_prompt,
                    config=genai_types.GenerateContentConfig(
                        response_mime_type="application/json",
                        response_schema=EvalScores,
                        temperature=0.1,
                    ),
                )
                scores = json.loads(response.text)
                all_scores.append(scores)
                logger.info(
                    f"  [{pos+1}/{len(valid_items)}] F={scores['faithfulness']:.2f}  "
                    f"AR={scores['answer_relevancy']:.2f}  "
                    f"CP={scores['context_precision']:.2f}  "
                    f"CR={scores['context_recall']:.2f}"
                )
                break  # Success, move to next question

            except Exception as e:
                err_str = str(e)
                if "429" in err_str or "RESOURCE_EXHAUSTED" in err_str:
                    # Extract retry delay from error message
                    wait_match = re.search(r"retry in (\d+\.?\d*)", err_str)
                    wait_secs = float(wait_match.group(1)) + 5 if wait_match else 65
                    wait_secs = min(wait_secs, 120)  # Cap at 2 minutes

                    if attempt < max_retries - 1:
                        logger.warning(
                            f"  ⏳ Rate limit câu {pos+1}/{len(valid_items)}, "
                            f"chờ {wait_secs:.0f}s rồi retry (attempt {attempt+2}/{max_retries})..."
                        )
                        time.sleep(wait_secs)
                    else:
                        logger.error(f"  ❌ Câu {pos+1}: Hết retry sau {max_retries} lần")
                else:
                    logger.error(f"  ❌ Lỗi đánh giá câu {pos+1}: {e}")
                    break  # Non-rate-limit error, skip this question

    if not all_scores:
        return {}

    # Tính trung bình
    avg_scores = {
        "faithfulness": sum(s["faithfulness"] for s in all_scores) / len(all_scores),
        "answer_relevancy": sum(s["answer_relevancy"] for s in all_scores) / len(all_scores),
        "context_precision": sum(s["context_precision"] for s in all_scores) / len(all_scores),
        "context_recall": sum(s["context_recall"] for s in all_scores) / len(all_scores),
    }

    logger.info(f"✅ Đánh giá hoàn thành! ({len(all_scores)}/{len(valid_items)} câu)")
    logger.info(f"  Điểm trung bình: {avg_scores}")
    return avg_scores


# =====================================================================
# MAIN - CHẠY TOÀN BỘ PIPELINE
# =====================================================================

async def main():
    """Pipeline chính: PDF → Q&A → RAG → Evaluate → Excel"""

    print("=" * 70)
    print("  ĐÁNH GIÁ HỆ THỐNG DocuGraph RAG BẰNG RAGAS METRICS")
    print("=" * 70)
    print()

    # ── Cấu hình ──
    PDF_PATH = os.path.join("..", "RAGS", "25-trang.pdf")
    NUM_QUESTIONS = 15   # Số câu hỏi cần sinh
    OUTPUT_EXCEL = os.path.join("..", "RAGS", f"ragas_evaluation_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")

    # Kiểm tra file PDF
    if not os.path.exists(PDF_PATH):
        logger.error(f"❌ Không tìm thấy file PDF: {PDF_PATH}")
        logger.info("Hãy đảm bảo file 25-trang.pdf nằm trong thư mục RAGS/")
        return

    # ── Bước 1: Trích xuất text ──
    print("\n📌 BƯỚC 1: Trích xuất text từ tài liệu PDF")
    print("-" * 50)
    doc_text = extract_pdf_text(PDF_PATH)
    if not doc_text.strip():
        logger.error("❌ Không trích xuất được text từ PDF!")
        return

    # ── Bước 2: Sinh bộ câu hỏi ──
    print("\n📌 BƯỚC 2: Sinh bộ câu hỏi & câu trả lời ground-truth")
    print("-" * 50)
    qa_pairs = generate_qa_pairs(doc_text, num_questions=NUM_QUESTIONS)
    if not qa_pairs:
        logger.error("❌ Không sinh được câu hỏi!")
        return

    # In ra bộ câu hỏi đã sinh
    print("\n📋 Bộ câu hỏi đã sinh:")
    for i, qa in enumerate(qa_pairs):
        print(f"  Q{i+1}: {qa['question']}")
        print(f"  A{i+1}: {qa['ground_truth'][:100]}...")
        print()

    # ── Bước 3: Chạy RAG ──
    print("\n📌 BƯỚC 3: Chạy RAG pipeline cho từng câu hỏi")
    print("-" * 50)

    # Tìm notebook_id từ Neo4j (lấy notebook đầu tiên có source)
    from app.services.neo4j_service import get_neo4j_driver
    driver = get_neo4j_driver()
    with driver.session() as session:
        result = session.run(
            "MATCH (n:Notebook)<-[:BELONGS_TO]-(s:Source)<-[:PART_OF]-(c:Chunk) "
            "RETURN DISTINCT n.id AS notebook_id, count(c) AS chunk_count "
            "ORDER BY chunk_count DESC LIMIT 1"
        )
        record = result.single()

    if not record:
        logger.error(
            "❌ Không tìm thấy notebook nào có dữ liệu trong Neo4j!\n"
            "   Hãy upload tài liệu vào hệ thống trước khi chạy đánh giá."
        )
        return

    notebook_id = record["notebook_id"]
    chunk_count = record["chunk_count"]
    logger.info(f"📓 Sử dụng notebook: {notebook_id} ({chunk_count} chunks)")

    rag_results = await run_rag_for_questions(qa_pairs, notebook_id)

    # ── Bước 4: Đánh giá ──
    print("\n📌 BƯỚC 4: Đánh giá bằng RAGAS metrics")
    print("-" * 50)

    ragas_scores = {}
    try:
        ragas_scores = evaluate_with_ragas(rag_results)
        logger.info("Sử dụng thư viện RAGAS chính thức.")
    except ImportError:
        logger.warning(
            "⚠️  Thư viện 'ragas' chưa được cài đặt. "
            "Sử dụng Gemini-based evaluation thay thế.\n"
            "   Để cài RAGAS: uv pip install ragas datasets"
        )
        ragas_scores = evaluate_without_ragas(rag_results)
    except Exception as e:
        logger.warning(f"⚠️  Lỗi khi chạy RAGAS: {e}. Dùng Gemini evaluation.")
        ragas_scores = evaluate_without_ragas(rag_results)

    # ── Bước 5: Xuất Excel ──
    print("\n📌 BƯỚC 5: Xuất kết quả ra file Excel")
    print("-" * 50)

    try:
        export_to_excel(rag_results, ragas_scores, OUTPUT_EXCEL)
    except ImportError:
        logger.error(
            "❌ Cần cài đặt openpyxl để xuất Excel:\n"
            "   uv pip install openpyxl"
        )
        return

    # ── Kết quả tổng hợp ──
    print("\n" + "=" * 70)
    print("  KẾT QUẢ ĐÁNH GIÁ TỔNG HỢP")
    print("=" * 70)
    print(f"  📊 Số câu hỏi đánh giá: {len(rag_results)}")
    for metric, score in (ragas_scores or {}).items():
        if isinstance(score, (int, float)):
            bar = "█" * int(score * 20) + "░" * (20 - int(score * 20))
            print(f"  {metric:25s}: {score:.4f} [{bar}]")
    print(f"\n  📁 File Excel: {os.path.abspath(OUTPUT_EXCEL)}")
    print("=" * 70)


if __name__ == "__main__":
    asyncio.run(main())
